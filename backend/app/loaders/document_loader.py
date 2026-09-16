import asyncio
import io
import re
from dataclasses import dataclass
from pathlib import Path


@dataclass(slots=True)
class ExtractedPage:
    page_number: int
    text: str
    heading: str | None = None
    used_ocr: bool = False


@dataclass(slots=True)
class ExtractionResult:
    pages: list[ExtractedPage]
    used_ocr: bool


def clean_text(value: str) -> str:
    value = value.replace("\x00", " ").replace("\r\n", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _extract(path: str, extension: str, ocr_languages: str, minimum_chars: int) -> ExtractionResult:
    if extension == ".pdf":
        import pymupdf
        import pytesseract
        from PIL import Image

        pages: list[ExtractedPage] = []
        used_ocr = False
        with pymupdf.open(path) as document:
            for index, page in enumerate(document):
                text = clean_text(page.get_text("text"))
                page_ocr = len(text) < minimum_chars
                if page_ocr:
                    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                    image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                    text = clean_text(pytesseract.image_to_string(image, lang=ocr_languages))
                    used_ocr = True
                pages.append(ExtractedPage(page_number=index + 1, text=text, used_ocr=page_ocr))
        return ExtractionResult(pages=pages, used_ocr=used_ocr)
    if extension in {".jpg", ".jpeg", ".png"}:
        import pytesseract
        from PIL import Image

        with Image.open(path) as image:
            text = clean_text(pytesseract.image_to_string(image, lang=ocr_languages))
        return ExtractionResult(pages=[ExtractedPage(page_number=1, text=text, used_ocr=True)], used_ocr=True)
    if extension == ".docx":
        from docx import Document

        document = Document(path)
        parts: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
        for table in document.tables:
            parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return ExtractionResult(pages=[ExtractedPage(page_number=1, text=clean_text("\n\n".join(parts)))], used_ocr=False)
    text = Path(path).read_text(encoding="utf-8", errors="strict")
    return ExtractionResult(pages=[ExtractedPage(page_number=1, text=clean_text(text))], used_ocr=False)


async def extract_document(path: str, extension: str, ocr_languages: str, minimum_chars: int) -> ExtractionResult:
    return await asyncio.to_thread(_extract, path, extension, ocr_languages, minimum_chars)
